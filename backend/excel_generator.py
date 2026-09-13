"""
Generate Excel output in the required 7-column format:
trg.naziv | kolicina | vrijednost | bruto | neto | tar.broj | zem.por
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


HEADERS = ["trg.naziv", "kolicina", "vrijednost", "bruto", "neto", "tar.broj", "zem.por"]
COL_WIDTHS = [40, 12, 14, 12, 12, 14, 10]

HEADER_BG = "1F4E79"   # dark blue
HEADER_FG = "FFFFFF"   # white
ALT_ROW_BG = "DCE6F1"  # light blue for alternating rows


def _thin_border():
    thin = Side(style="thin")
    return Border(left=thin, right=thin, top=thin, bottom=thin)


def generate_excel(items: list[dict], output_path: str, consignee: str = "") -> str:
    """
    Generate Excel file from list of item dicts.

    Each item dict must have:
        name, quantity, value, gross_weight, net_weight, tariff_code, country_origin

    Returns the output file path.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Deklaracija"

    # Optional consignee row at top
    if consignee:
        ws.merge_cells(f"A1:{get_column_letter(len(HEADERS))}1")
        cell = ws["A1"]
        cell.value = consignee
        cell.font = Font(bold=True, size=12)
        cell.alignment = Alignment(horizontal="center")
        start_row = 2
    else:
        start_row = 1

    # Header row
    for col_idx, (header, width) in enumerate(zip(HEADERS, COL_WIDTHS), start=1):
        cell = ws.cell(row=start_row, column=col_idx, value=header)
        cell.font = Font(bold=True, color=HEADER_FG, size=10)
        cell.fill = PatternFill("solid", fgColor=HEADER_BG)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = _thin_border()
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.row_dimensions[start_row].height = 18

    # Data rows
    for row_offset, item in enumerate(items):
        row_num = start_row + 1 + row_offset
        is_alt = (row_offset % 2 == 1)
        bg_color = ALT_ROW_BG if is_alt else "FFFFFF"

        values = [
            item.get("name", ""),
            item.get("quantity", 0),
            item.get("value", 0),
            item.get("gross_weight", 0),
            item.get("net_weight", 0),
            item.get("tariff_code", ""),
            item.get("country_origin", ""),
        ]

        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row_num, column=col_idx, value=value)
            cell.fill = PatternFill("solid", fgColor=bg_color)
            cell.border = _thin_border()

            # Numeric columns: right-align and format
            if col_idx in (2, 3, 4, 5):  # qty, value, gross, net
                cell.alignment = Alignment(horizontal="right")
                if isinstance(value, (int, float)) and value != 0:
                    cell.number_format = "#,##0.00"
            else:
                cell.alignment = Alignment(horizontal="left")

    # Totals row
    if items:
        total_row = start_row + 1 + len(items)
        ws.cell(row=total_row, column=1, value="UKUPNO").font = Font(bold=True)

        total_qty = sum(item.get("quantity", 0) for item in items)
        total_val = sum(item.get("value", 0) for item in items)
        total_gross = sum(item.get("gross_weight", 0) for item in items)
        total_net = sum(item.get("net_weight", 0) for item in items)

        for col_idx, total in [(2, total_qty), (3, total_val), (4, total_gross), (5, total_net)]:
            cell = ws.cell(row=total_row, column=col_idx, value=total)
            cell.font = Font(bold=True)
            cell.number_format = "#,##0.00"
            cell.alignment = Alignment(horizontal="right")
            cell.border = _thin_border()

        for col_idx in range(1, len(HEADERS) + 1):
            ws.cell(row=total_row, column=col_idx).fill = PatternFill("solid", fgColor="BDD7EE")

    wb.save(output_path)
    return output_path
